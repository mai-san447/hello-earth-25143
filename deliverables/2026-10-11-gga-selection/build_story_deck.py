"""GGA セレクション（5分）用、ストーリーテリング版のスライドをつくる。

セリフの一文ごとに一枚、大きな言葉を一つだけ置く。スライドは読ませるものではなく、話の後ろで支えるもの。
セリフは本人がまとめた5分台本（2026-10-13）。ノートに、その一枚で話す言葉と時間を入れる。
9枚版（build_deck.py）は予備として残す。

使い方: python build_story_deck.py  →  同じフォルダに MORUNE25143_GGAセレクション_ストーリー_2026-10-13.pptx
"""
import random

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches

from deck_common import (HERE, IMG, NIGHT, PANEL, GOLD, TEXT, MUTED, DIM, SERIF,
                         text, panel, picture, dot)

OUT = HERE / "MORUNE25143_GGAセレクション_ストーリー_2026-10-13.pptx"
W, H = 13.333, 7.5

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]
count = 0


def sky(slide, seed):
    """夜空の地に、星を少しだけ散らす。毎回同じ配置になるよう seed を固定する"""
    rnd = random.Random(seed)
    for _ in range(26):
        d = rnd.choice([0.03, 0.03, 0.04, 0.05])
        c = rnd.choice([DIM, DIM, MUTED])
        dot(slide, rnd.uniform(0.2, W - 0.2), rnd.uniform(0.2, H - 0.2), d, c)


def slide(chapter, notes, seed=None):
    """chapter: 左上の小さな章名（例 "01 聞いてきた声"）。notes: その一枚で話す言葉"""
    global count
    count += 1
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NIGHT
    sky(s, seed if seed is not None else count)
    if chapter:
        text(s, 0.6, 0.4, 8, 0.4, chapter, size=12, color=GOLD)
    s.notes_slide.notes_text_frame.text = notes
    return s


def statement(chapter, notes, lines, size=40, color=TEXT, sub=None, y=2.3):
    """大きな言葉を一つだけ置く一枚。lines は行のリスト（行の頭に来ない位置で自分で改行する）"""
    s = slide(chapter, notes)
    text(s, 0.9, y, W - 1.8, 3.2, lines, size=size, font=SERIF, color=color, spacing=1.35)
    if sub:
        text(s, 0.9, 6.2, W - 1.8, 0.8, sub, size=14, color=MUTED)
    return s


# ---- 表紙・出場目的 ----------------------------------------------------------------
s = slide("", "【0:00】\nMORUNE 25143の中村です。\n［送る］", seed=101)
text(s, 0.9, 1.0, 7, 0.4, "LEAVE YOUR WISH WITH THE STARS", size=14, color=GOLD, font="Cormorant Garamond")
text(s, 0.9, 1.6, 7.4, 2.6, ["やりたいことを", "星に預けておく"], size=54, font=SERIF, spacing=1.15)
text(s, 0.9, 4.3, 7, 0.6, [[("MORUNE ", {"bold": True}), ("25143", {"color": GOLD, "bold": True})]], size=26)
text(s, 0.9, 6.5, 7, 0.4, "中村 舞衣子　G's ACADEMY 卒業制作", size=12, color=MUTED)
picture(s, "max-home.png", 9.3, 0.35, h=6.8)

s = slide("出場目的", "【0:03】\n今日の出場目的は、起業です。\n実装と検証を一緒に進める仲間を探しています。\n［送る］")
text(s, 0.9, 2.2, W - 1.8, 1.2, "起業", size=60, font=SERIF, color=GOLD)
text(s, 0.9, 3.7, W - 1.8, 1.0, "実装と検証を、一緒に進める仲間を探しています", size=30, font=SERIF)

# ---- 01 聞いてきた声 ----------------------------------------------------------------
s = slide("01 聞いてきた声", "【0:10】\n私は調査・研究を行うシンクタンクでマーケティングを担当し、副業では働く女性の社外メンターをしています。\n［送る］")
text(s, 0.9, 2.2, W - 1.8, 1.0, "中村 舞衣子", size=44, font=SERIF)
text(s, 0.9, 3.5, W - 1.8, 1.6, ["調査・研究を行うシンクタンクで、マーケティングを担当", "副業で、働く女性の社外メンター", "大学院で、省察支援を研究（JSET 2026）"], size=20, color=MUTED, spacing=1.6)

statement("01 聞いてきた声", "【0:20】\n面談で最も多く聞いてきたのは、仕事と家庭の両立、つまり時間の制約に関する悩みでした。\n［送る］",
          [[("面談で、最も多く聞いてきた悩み", {"size": 20, "color": MUTED, "font": "メイリオ"})], "仕事と家庭の両立。", [("つまり、時間の制約。", {"color": GOLD})]], size=44, y=2.0)

# ---- 02 意欲を止める、時間の制約 ------------------------------------------------------
statement("02 意欲を止める、時間の制約", "【0:40】\n聞いたのは、「やりたいことがない」という声ではありません。\n［送る］",
          ["「やりたいことがない」", [("という声では、ない。", {"color": MUTED})]], size=44)
statement("02 意欲を止める、時間の制約", "【0:45】\n「活躍したい」「学び直したい」。\n［送る］",
          ["「活躍したい」", "「学び直したい」"], size=48)
statement("02 意欲を止める、時間の制約", "【0:48】\nでも時間がなく、自信まで失っていく。\n［送る］",
          ["でも、時間がない。", [("自信まで、失っていく。", {"color": GOLD})]], size=48)
s = slide("02 意欲を止める、時間の制約", "【0:52】\nMORUNE 25143は、育児や介護、療養などで願いをいったん手放した人が、急かされずに預け、自分のタイミングで受け取り直す体験です。\n［送る］")
text(s, 0.9, 1.4, 7.6, 0.5, "MORUNE 25143", size=18, color=GOLD, bold=True)
text(s, 0.9, 2.1, 7.8, 3.6, ["育児や介護、療養などで", "願いをいったん手放した人が", "急かされずに預け、", [("自分のタイミングで受け取り直す", {"color": GOLD})]], size=30, font=SERIF, spacing=1.45)
text(s, 0.9, 6.3, 7.8, 0.6, "参考：6歳未満の子どもがいる世帯の家事関連時間は1日あたり妻7時間28分・夫1時間54分（総務省「令和3年社会生活基本調査」）", size=10, color=DIM)
picture(s, "02-deposit.png", 9.4, 0.5, h=6.5)

# ---- 03 研究から得た設計の足場 --------------------------------------------------------
statement("03 研究から得た設計の足場", "【1:00】\n「いつかやりたい」は、ToDoにすると重荷になります。\n［送る］",
          ["「いつかやりたい」は、", [("ToDo にすると、重荷になる。", {"color": MUTED})]], size=44)
statement("03 研究から得た設計の足場", "【1:05】\n大学院で、止まったふりかえりの再開を研究して分かったのは、必要なのは強い決意ではなく、小さなきっかけだということ。\n［送る］",
          [[("必要なのは、強い決意ではなく", {"color": MUTED})], [("小さなきっかけ。", {"color": GOLD})]], size=48,
          sub="大学院の研究：省察の再起動を支援する軽量学習ログシステムの設計要件（JSET 2026）")
s = slide("03 研究から得た設計の足場", "【1:12】\n手がかりは小さく、仕組みは軽く、迷ったまま選べて、つながりを感じられる。\nこの四つを体験の設計にしました。\n［送る］")
conds = [("手がかりは小さく", "願いは12字だけ"),
         ("仕組みは軽く", "預けるだけ。期限も通知もない"),
         ("迷ったまま選べる", "受け取る／もう少し預ける／手放す"),
         ("つながりを感じられる", "みんなの星空・支援の場")]
text(s, 0.9, 1.2, 6, 0.4, "四つの足場", size=14, color=GOLD, bold=True)
text(s, 7.5, 1.2, 5, 0.4, "MORUNE では", size=14, color=GOLD, bold=True)
for i, (h, m) in enumerate(conds):
    y = 1.75 + i * 1.25
    text(s, 0.9, y, 6.2, 1.0, h, size=30, font=SERIF, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 6.75, y, 0.6, 1.0, "→", size=26, color=GOLD, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 7.5, y, 5.4, 1.0, m, size=20, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)

# ---- 04 はやぶさ ------------------------------------------------------------------
statement("04 はやぶさ", "【1:20】\nこの体験に重ねたのが、はやぶさの帰還です。\n［送る］",
          ["この体験に重ねたのは、", [("はやぶさの帰還。", {"color": GOLD})]], size=48)
s = slide("04 はやぶさ", "【1:25】\n編集長がキャリアブレイクに入り、私が編集長代理を務めた時期に、川口淳一郎先生を取材しました。\n「高い塔を建てなければ、新たな水平線は見えてこない」。新しい役割へ踏み出す私を支えた言葉です。\n［送る］")
text(s, 0.9, 1.6, W - 1.8, 0.6, "編集長代理を務めた時期に、川口淳一郎先生（はやぶさ プロジェクトマネージャー）を取材", size=16, color=MUTED)
text(s, 0.9, 2.7, W - 1.8, 2.2, ["「高い塔を建てなければ、", "新たな水平線は見えてこない」"], size=44, font=SERIF, color=GOLD, spacing=1.35)
text(s, 0.9, 5.6, W - 1.8, 0.6, "新しい役割へ踏み出す私を、支えた言葉", size=18, color=TEXT)
s = slide("04 はやぶさ", "【1:38】\n手放した願いを遠い星に預け、カプセルのように一つずつ帰す。これが作品のメタファーです。\n（間）\n見てください。\n［送る・動画を再生］")
text(s, 0.9, 2.0, 7.6, 3.4, ["手放した願いを、", "遠い星に預ける。", [("カプセルのように、", {"color": GOLD})], [("一つずつ帰ってくる。", {"color": GOLD})]], size=36, font=SERIF, spacing=1.45)
picture(s, "06-landed.png", 9.4, 0.5, h=6.5)

# ---- 05 デモ ----------------------------------------------------------------------
s = slide("05 デモ", """【1:45–3:47】動画 2分02秒。［ ］の時間は動画の中の時間
［0:00］仕事と家のことに追われ、「ギターをもう一度」を打ち消してきた人がいます。ここは、今いる場所、今の時刻の北の空です。
［0:19］願いを書くと、星が灯ります。期限も通知もありません。「今はできなくても、忘れなくていい」と、少し心が軽くなります。
［0:43］余白ができた日に星へ触れると、カプセルが帰ってきます。「まだ、やってみたい」と気づく瞬間です。
［0:58］受け取る、もう少し預ける、手放す。今の自分に合う距離を選べます。
［1:09］受け取ると帰還票になります。書くのは「日程を調べる」など、最初の一歩だけ。
［帰還票とキーホルダーまたはブレスレットを見せる］持ち帰れば、願いを日常で思い出す目印になります。
［1:36］みんなの星空では、内容を見せず、ほかの人の願いを光として感じられます。「立ち止まっているのは自分だけではない」。そのつながりが次の一歩を支えます。
［動画終了後に送る］""")
s.shapes.add_movie(str(HERE / "demo-16x9.mp4"), Inches(1.6), Inches(1.0), Inches(10.1), Inches(5.68),
                   poster_frame_image=str(IMG / "demo-poster.png"), mime_type="video/mp4")
text(s, 1.6, 6.8, 10.1, 0.3, "▶ クリックで再生・2分02秒（このPCの中だけのデータで撮影。見本の願い入り）", size=10, color=MUTED, align=PP_ALIGN.CENTER)

# ---- 06 技術 ----------------------------------------------------------------------
s = slide("06 技術", "【3:47】\nコードはAIに書かせ、何を作るかは私が決め、数で確かめました。\n［送る］")
text(s, 0.9, 1.5, W - 1.8, 1.8, ["AI に書かせ、", [("何を作るかは、わたしが決め、数で確かめた。", {"color": GOLD})]], size=34, font=SERIF, spacing=1.4)
text(s, 0.9, 3.9, 6, 0.4, "星が重なる日（1年あたり）", size=14, color=MUTED)
text(s, 0.9, 4.4, 11, 1.4, [[("360日", {"color": MUTED}), ("　→　", {"color": GOLD}), ("0日", {"color": GOLD})]], size=60, font=SERIF)
text(s, 0.9, 6.1, W - 1.8, 0.6, "作って試してから、やめたもの：AI で願いの絵をつくる機能・星座の線・応援の信号", size=13, color=DIM)
statement("06 技術", "【3:53】\n位置情報は外へ送らず、オフラインでも動きます。\n願いは初めは端末の中だけ。「みんなの星空に出す」と選んだものだけを送り、イトカワまでの距離も計算しています。\n［送る］",
          ["位置情報は、外へ送らない。", [("願いは、初めは端末の中だけ。", {"color": GOLD})]], size=40,
          sub="オフラインでも動く ／ 「みんなの星空に出す」と選んだ願いだけを送る ／ イトカワまでの距離も計算")

# ---- 07 届け方と事業モデル --------------------------------------------------------------
statement("07 届け方と事業モデル", "【4:02】\n入口は、展示とキャリア・復職支援の場です。\n［送る］",
          ["入口は、", [("展示と、キャリア・復職支援の場。", {"color": GOLD})]], size=44)
s = slide("07 届け方と事業モデル", "【4:06】\n個人利用は無料。クライアントへの導入と、イベントでの帰還票、キーホルダー、ブレスレットの物販で続けます。\nこれらは、デジタルの願いを日常へつなぐ目印です。\n［送る］")
text(s, 0.9, 1.4, 7.6, 1.0, "個人は、無料。", size=44, font=SERIF)
text(s, 0.9, 2.7, 7.6, 2.4, ["クライアントへの導入と、", "日常へつなぐ目印の物販で続ける。"], size=26, font=SERIF, color=GOLD, spacing=1.45)
text(s, 0.9, 5.2, 7.6, 1.0, "帰還票 ／ キーホルダー ／ ブレスレット", size=18, color=MUTED)
picture(s, "10b-receipt-step.png", 9.4, 0.5, h=6.5)

# ---- 08 今後の検証とお願い --------------------------------------------------------------
statement("08 今後の検証とお願い", "【4:17】\nこの作品を作る途中で、私も立ち止まる側になりました。\n（間）\n［送る］",
          ["この作品を作る途中で、", "わたしも、立ち止まる側になった。"], size=40)
statement("08 今後の検証とお願い", "【4:25】\nだからこれは、願いを急かさず、「もう一度」と思えた日に返す仕組みです。\n［送る］",
          ["願いを急かさず、", [("「もう一度」と思えた日に、返す。", {"color": GOLD})]], size=44)
statement("08 今後の検証とお願い", "【4:32】\nまず、展示と支援の場で検証します。\n［送る］",
          ["まず、", [("展示と支援の場で、検証する。", {"color": GOLD})]], size=44)
s = slide("08 今後の検証とお願い", "【4:36】\n展示の機会、検証の場、実装を共にする仲間を探しています。\nQRから、願いを一つ預けられます。\nありがとうございました。")
text(s, 0.9, 1.2, 8.2, 1.0, "やりたかったことを、置き去りにしない", size=30, font=SERIF)
asks = [("展示の機会", "科学館・プラネタリウム・自治体のイベント"),
        ("検証の場", "キャリアや復職を支える人事・支援者のもとで"),
        ("実装の仲間", "実装と検証を一緒に進めるエンジニア")]
for i, (h, b) in enumerate(asks):
    y = 2.5 + i * 1.25
    panel(s, 0.9, y, 8.0, 1.05, line=GOLD if i == 2 else None)
    text(s, 1.2, y, 2.4, 1.05, h, size=22, font=SERIF, color=GOLD, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 3.6, y, 5.2, 1.05, b, size=15, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
panel(s, 9.6, 2.3, 3.2, 4.1, fill=RGBColor(0xFF, 0xFF, 0xFF))
picture(s, "qr-site.png", 9.85, 2.5, w=2.7)
text(s, 9.6, 5.3, 3.2, 1.0, ["いま、ひとつ預けられます", "morune-25143.morune-25143.workers.dev"], size=10, color=NIGHT, align=PP_ALIGN.CENTER)

prs.save(OUT)
print("saved", OUT, count, "slides")
