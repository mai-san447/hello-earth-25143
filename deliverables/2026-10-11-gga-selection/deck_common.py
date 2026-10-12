"""セレクション用スライドの共通部品（色・字・図形）。build_story_deck.py から使う。

build_deck.py（9枚版）は同じ部品を自前で持っている。9枚版は予備として凍結しているので、
二重になっているのは承知のうえ（直すなら build_deck.py もこちらを使う形に寄せる）。
"""
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
IMG = HERE / "img"

# アプリの画面と同じ色（public/style.css の夜空・金色）にそろえる
NIGHT = RGBColor(0x0B, 0x10, 0x20)
PANEL = RGBColor(0x16, 0x1D, 0x30)
GOLD = RGBColor(0xE2, 0xB8, 0x67)
TEXT = RGBColor(0xF2, 0xEE, 0xE6)
MUTED = RGBColor(0x9A, 0xA3, 0xB5)
DIM = RGBColor(0x5E, 0x67, 0x7D)

SERIF = "游明朝 Demibold"  # 見出し。アプリの明朝に近い字形で、どの Windows にもある
SANS = "メイリオ"          # 本文。python-pptx では ea/latin を明示しないと置き換わる


def set_font(run, size, color=TEXT, bold=False, font=SANS):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", font)


def text(slide, x, y, w, h, lines, size=18, color=TEXT, bold=False, font=SANS,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.25):
    """lines: 文字列、または段落のリスト。段落は文字列か [(文字, {上書き}), ...]"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    if isinstance(lines, str):
        lines = [lines]
    for i, para in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        runs = [(para, {})] if isinstance(para, str) else para
        for s, opt in runs:
            r = p.add_run()
            r.text = s
            set_font(r, opt.get("size", size), opt.get("color", color),
                     opt.get("bold", bold), opt.get("font", font))
    return box


def panel(slide, x, y, w, h, fill=PANEL, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.25)
    shp.shadow.inherit = False
    return shp


def picture(slide, name, x, y, h=None, w=None):
    kw = {}
    if h is not None:
        kw["height"] = Inches(h)
    if w is not None:
        kw["width"] = Inches(w)
    return slide.shapes.add_picture(str(IMG / name), Inches(x), Inches(y), **kw)


def dot(slide, x, y, d, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp
