"""GGA セレクション（10分）用スライドの下書きをつくる。

中身の出どころは docs/GGAセレクション企画書.html と docs/ペルソナ.md。
仕上げ（言い回し・写真の差し替え）は PowerPoint で本人が行う前提なので、
凝った装飾はせず、1枚1メッセージと画面の大きさを優先する。

使い方: python build_deck.py  →  同じフォルダに MORUNE25143_GGAセレクション_2026-10-12.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
IMG = HERE / "img"
OUT = HERE / "MORUNE25143_GGAセレクション_2026-10-12.pptx"

# アプリの画面と同じ色（public/style.css の夜空・金色）にそろえる
NIGHT = RGBColor(0x0B, 0x10, 0x20)
PANEL = RGBColor(0x16, 0x1D, 0x30)
GOLD = RGBColor(0xE2, 0xB8, 0x67)
TEXT = RGBColor(0xF2, 0xEE, 0xE6)
MUTED = RGBColor(0x9A, 0xA3, 0xB5)

SERIF = "游明朝 Demibold"  # 見出し。アプリの明朝に近い字形で、どの Windows にもある
SANS = "メイリオ"          # 本文。python-pptx では ea/latin を明示しないと置き換わる

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def set_font(run, size, color=TEXT, bold=False, font=SANS):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", font)


def text(slide, x, y, w, h, lines, size=18, color=TEXT, bold=False, font=SANS,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.25):
    """lines: 文字列、または段落のリスト。段落は文字列か [(文字, {上書き}), ...]"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    if isinstance(lines, str):
        lines = [lines]
    for i, para in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        runs = [(para, {})] if isinstance(para, str) else para
        for s, opt in runs:
            r = p.add_run()
            r.text = s
            set_font(r, opt.get("size", size), opt.get("color", color),
                     opt.get("bold", bold), opt.get("font", font))
    return box


def panel(slide, x, y, w, h, fill=PANEL, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.25)
    shp.shadow.inherit = False
    return shp


def picture(slide, name, x, y, h=None, w=None):
    kw = {}
    if h is not None:
        kw["height"] = Inches(h)
    if w is not None:
        kw["width"] = Inches(w)
    return slide.shapes.add_picture(str(IMG / name), Inches(x), Inches(y), **kw)


def new_slide(no, label, title, notes, title_size=30):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NIGHT
    if no:
        text(s, 0.6, 0.35, 8, 0.4, [[(no + "  ", {"color": GOLD, "bold": True}), (label, {"color": MUTED})]], size=12)
    if title:
        text(s, 0.6, 0.7, 12.1, 1.2, title, size=title_size, font=SERIF)
    s.notes_slide.notes_text_frame.text = notes
    return s



# GGA セレクションの持ち時間10分は「発表＋講評＋入れ替え」。発表は5分強に収める（2026-10-12、説明会の資料から）。
# 最初に出場目的を宣言し、デモを必ず入れる。9枚・約5分40秒。

def flow_arrow(slide, x, y):
    text(slide, x, y, 0.5, 0.6, "→", size=26, color=GOLD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# 0 表紙・出場目的 --------------------------------------------------------------
s = new_slide("", "", "", """【0:00–0:15】出場目的を最初に宣言する
・MORUNE 25143 の中村です。
・今日の出場目的は、起業です。この体験を、展示と、働く人を支える場で届ける仲間を探しています。""")
text(s, 0.8, 0.9, 7, 0.4, "LEAVE YOUR WISH WITH THE STARS", size=14, color=GOLD, font="Cormorant Garamond")
text(s, 0.8, 1.4, 7.2, 2.4, ["やりたいことを", "星に預けておく"], size=48, font=SERIF, spacing=1.15)
text(s, 0.8, 3.65, 7, 0.6, [[("MORUNE ", {"bold": True}), ("25143", {"color": GOLD, "bold": True})]], size=24)
panel(s, 0.8, 4.6, 7.2, 1.5, line=GOLD)
text(s, 1.1, 4.7, 6.7, 0.45, "出場目的", size=13, color=GOLD, bold=True)
text(s, 1.1, 5.1, 6.7, 0.95, ["起業（スモールビジネス）", "この体験を、展示と支援の場で届ける仲間を探しています"], size=18, spacing=1.35)
text(s, 0.8, 6.55, 7, 0.4, "中村 舞衣子　G's ACADEMY 卒業制作", size=12, color=MUTED)
picture(s, "max-home.png", 9.3, 0.35, h=6.8)

# 1 課題 ------------------------------------------------------------------------
s = new_slide("01", "Problem", "「今じゃなくていい」のまま、やりたかったことを忘れてしまう", """【0:15–0:55】
・療養や休職、育休や介護。自分を後回しにする時期があります。
・その間にも「いつかやりたいこと」は浮かびます。でも今じゃない、と思って書かない。
・ToDo に入れると重荷になる。SNS に書くのは違う。メモは見返さない。置き場所がないまま、忘れてしまいます。
・戻ったときには思い出せないか、たまりすぎて選べない。""", title_size=25)
steps = [("後回しにする時期", "療養・休職・育休・介護\n忙しい時期"),
         ("「今じゃなくていい」", "やりたいことは浮かぶが\nToDo は重荷、SNS は違う\nメモは見返さない"),
         ("忘れてしまう", "戻ったときに思い出せない\nたまりすぎて選べない")]
for i, (h, b) in enumerate(steps):
    x = 0.6 + i * 4.2
    panel(s, x, 2.2, 3.6, 3.2, line=GOLD if i == 2 else None)
    text(s, x + 0.3, 2.45, 3.0, 0.6, h, size=20, font=SERIF, color=GOLD if i == 2 else TEXT)
    text(s, x + 0.3, 3.25, 3.0, 2.0, b, size=15, color=MUTED, spacing=1.5)
    if i < 2:
        flow_arrow(s, x + 3.62, 3.4)
text(s, 0.6, 5.85, 12.2, 0.9, "最初に届ける人：しばらく仕事を離れる人（仮の人物「みわさん 42歳」）。休みの前の晩、「天文台で土星を見たい」と思うが、どこにも書かない", size=14, color=MUTED)

# 2 なぜわたしか -----------------------------------------------------------------
s = new_slide("02", "Why me", "なぜ、わたしが願いを宇宙に預けるのか", """【0:55–1:25】
・前職で、はやぶさのプロジェクトマネージャーを取材しました。はやぶさは7年かけて、小惑星 25143 イトカワから帰ってきました。作品名の 25143 はこの番号です。
・社外の女性メンターとして話を聞く中で、自分のキャリアを描けない人に出会ってきました。
・昔やりたかったこと、好きなことが置き去りになっている。自分を後回しにしてきた人たちです。
・この作品は、わたしがメンターとしてしている支援の、1つの形です。
・大学院では 1on1 のふりかえりを研究し、過去の自分の言葉に出会い直すと、次の一歩を選びやすくなると学びました。""")
cards = [("取材", "はやぶさは7年かけて\n小惑星 25143 イトカワから\n帰ってきた\n作品名は、この番号から"),
         ("社外女性メンター", "自分のキャリアを\n描けない人に出会ってきた\n昔やりたかったこと\n好きなことが置き去り\n自分を後回しにしてきた"),
         ("研究", "過去の自分の言葉に\n出会い直すと\n次の一歩を選びやすくなる\n（大学院・JSET 2026）")]
for i, (k, b) in enumerate(cards):
    x = 0.6 + i * 4.15
    panel(s, x, 2.0, 3.85, 3.8)
    text(s, x + 0.3, 2.25, 3.2, 0.6, k, size=24, font=SERIF, color=GOLD)
    text(s, x + 0.3, 3.05, 3.25, 2.6, b, size=15, color=MUTED, spacing=1.5)

# 3 解決 ------------------------------------------------------------------------
s = new_slide("03", "Solution", "忘れないように星に預け、やりたくなったら自分で呼び戻す", """【1:25–1:55】
・やりたいことを12字で星に預けます。本物の星空の、北の星座の星に灯ります。19個まで。
・何もしなくても星は暗くなりません。期限も通知もありません。
・やりたくなったら、その星に触れると、願いが帰ってきます。
・受け取ると帰還票ができて、最初の小さな一歩を1つ書きます。""", title_size=27)
flow = [("預ける", "02-deposit.png", "12字で、星に預ける"),
        ("待つ", "max-home.png", "本物の星空に灯る\n期限も通知もない"),
        ("呼び戻す", "08-card.png", "やりたくなったら\n星に触れる"),
        ("一歩を書く", "10b-receipt-step.png", "帰還票に\n最初の小さな一歩")]
for i, (h, img, b) in enumerate(flow):
    x = 0.55 + i * 3.2
    picture(s, img, x + 0.35, 1.85, h=3.6)
    text(s, x, 5.55, 2.6, 0.5, h, size=19, font=SERIF, align=PP_ALIGN.CENTER, color=GOLD)
    text(s, x, 6.05, 2.6, 0.9, b, size=12, color=MUTED, align=PP_ALIGN.CENTER)
    if i < 3:
        flow_arrow(s, x + 2.6, 3.35)

# 4 デモ ------------------------------------------------------------------------
s = new_slide("04", "Demo", "デモ：30日前に預けた願いを、呼び戻す", """【1:55–3:15】動画を再生する（約70秒）
・預けます。「ギターをもう一度」。初めは、わたしの星空にだけ灯ります。
・30日前に預けた「天文台で土星の輪を見る」の星に触れます。カプセルが地球に着きます。
・「30日ぶりに帰ってきました」。受け取るか、もう少し預けるか、手放すか。
・受け取ると、帰還票が現像されます。最初の小さな一歩は「土曜の観望会を調べる」。
・最後に、みんなの星空を遠くから眺めます。ほかの人の願いが、光として灯っています。""")
s.shapes.add_movie(str(HERE / "demo-16x9.mp4"), Inches(1.6), Inches(1.55), Inches(10.1), Inches(5.68),
                   poster_frame_image=str(IMG / "demo-poster.png"), mime_type="video/mp4")
text(s, 1.6, 7.15, 10.1, 0.3, "▶ クリックで再生（このPCの中だけのデータで撮影。見本の願い入り）", size=10, color=MUTED, align=PP_ALIGN.CENTER)

# 5 急かさない・比べない・晒さない ---------------------------------------------------
s = new_slide("05", "Care", "急かさない。比べない。晒さない。", """【3:15–3:55】
・急かさない：期限も、通知も、連続日数もありません。何もしなくても星は暗くなりません。今でなければ「もう少し預ける」、要らなくなれば「手放す」。
・比べない：人と比べる数を持ちません。
・晒さない：願いは、初めは自分の星空だけ。みんなの星空では、ほかの人の言葉ではなく光を分け合います。「自分だけじゃない」は感じられて、傷つく言葉は入りません。つらい言葉を書いた人には、端末の中だけで相談先を案内します。""")
cols = [("急かさない", "期限・通知・連続日数なし\n何もしなくても\n星は暗くならない\n今でなければ\n「もう少し預ける」"),
        ("比べない", "人と比べる数を持たない\n一歩ごとに\n自分の星だけが\n明るくなる"),
        ("晒さない", "初めは\nわたしの星空だけ\nみんなの星空は\n言葉ではなく光\nつらい言葉には\n相談先を案内")]
for i, (h, b) in enumerate(cols):
    x = 0.6 + i * 3.05
    panel(s, x, 1.95, 2.85, 4.9)
    text(s, x + 0.25, 2.2, 2.4, 0.6, h, size=22, font=SERIF, color=GOLD)
    text(s, x + 0.25, 3.0, 2.45, 3.7, b, size=13, color=MUTED, spacing=1.7)
picture(s, "12-everyone.png", 10.0, 1.95, h=4.9)
text(s, 9.55, 6.9, 3.4, 0.4, "みんなの星空：ほかの人の願いは光だけ", size=11, color=MUTED, align=PP_ALIGN.CENTER)

# 6 技術の主体性 ------------------------------------------------------------------
s = new_slide("06", "Technology", "AI に書かせ、何を作るかは自分で決め、数で確かめた", """【3:55–4:35】
・コードは Claude と Codex に書かせ、互いにレビューさせました。何を作って何をやめるかは、わたしが決めました。AI で願いの絵をつくる機能、星座の線、応援の信号は、作って試してからやめました。
・決めるときは数で確かめました。願いを30個まで置ける作りでは、1年のうち360日ほど、星が重なる日がありました。19の席にして0日にしました。同じ願いが続けて帰る確率も、50%から0%にしました。
・本物の星空は、5,080個の星を時刻と場所から計算しています。テストは238件です。""", title_size=26)
text(s, 0.6, 1.9, 5.8, 0.4, "作ってから、やめたもの", size=14, color=GOLD, bold=True)
quits = [("AI で願いの絵をつくる", "願いのイメージとずれた"), ("星座の線", "意味が伝わらなかった"), ("応援の信号", "シェアとの違いが伝わらなかった"), ("「帰す」ボタン", "触れれば帰る形にした")]
for i, (a, b) in enumerate(quits):
    y = 2.35 + i * 0.85
    panel(s, 0.6, y, 5.8, 0.72)
    text(s, 0.8, y, 2.9, 0.72, a, size=14, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 3.7, y, 2.6, 0.72, b, size=12, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 5.85, 5.8, 1.0, "本物の星空：5,080個の星を時刻と場所から計算\n自動テスト 238件", size=13, color=MUTED, spacing=1.5)
text(s, 7.0, 1.9, 5.8, 0.4, "数で確かめてから決めた（シミュレーション）", size=14, color=GOLD, bold=True)

def bars(slide, y, label, before, after, unit, maximum):
    text(slide, 7.0, y, 5.8, 0.4, label, size=13)
    for k, (name, value, color) in enumerate([("前", before, MUTED), ("今", after, GOLD)]):
        yy = y + 0.45 + k * 0.5
        text(slide, 7.0, yy, 0.5, 0.4, name, size=12, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
        width = max(0.05, 4.2 * value / maximum)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.5), Inches(yy + 0.08), Inches(width), Inches(0.26))
        bar.fill.solid(); bar.fill.fore_color.rgb = color; bar.line.fill.background(); bar.shadow.inherit = False
        text(slide, 7.6 + width, yy, 1.6, 0.4, f"{value}{unit}", size=13, color=color, anchor=MSO_ANCHOR.MIDDLE)
bars(s, 2.4, "星が重なる日（1年あたり、まとめて預ける人）", 363, 0, "日", 365)
bars(s, 3.95, "同じ願いが続けて帰る確率（候補が2つのとき）", 50, 0, "%", 100)
text(s, 7.0, 5.5, 5.8, 1.3, "旧：番号で星の場所を決め、上限30個\n今：空いている席に置き、上限19個（北の星座3つの星の数）", size=12, color=MUTED, spacing=1.5)

# 7 届け方と仕組み ----------------------------------------------------------------
s = new_slide("07", "Distribution", "展示で出会い、紙で持ち帰り、アプリで続く", """【4:35–5:15】
・最初の入口は展示です。科学館や宇宙イベントで、QR から願いを預け、その場で帰還票を印刷して渡します。キャリアや復職を支える場でも、同じ形で使えます。
・紙とブレスレットが、家まで願いを連れて帰ります。帰還票の QR から、自分の星に戻れます。
・個人からはお金を取りません。展示の導入、持ち帰る物、アーティストとの共創で成り立たせます。""", title_size=28)
road = [("展示・支援の場", "科学館・宇宙イベント\nキャリア・復職の支援の場\nQR から預ける\n帰還票をその場で印刷"),
        ("紙と物", "帰還票・ブレスレット\n最初の一歩が目に入る\nQR から自分の星へ"),
        ("アプリ", "やりたくなったら\n星に触れて呼び戻す")]
for i, (h, b) in enumerate(road):
    x = 0.6 + i * 4.2
    panel(s, x, 1.95, 3.6, 2.65, line=GOLD if i == 0 else None)
    text(s, x + 0.3, 2.15, 3.0, 0.6, h, size=22, font=SERIF, color=GOLD)
    text(s, x + 0.3, 2.85, 3.0, 1.7, b, size=14, color=MUTED, spacing=1.5)
    if i < 2:
        flow_arrow(s, x + 3.62, 2.95)
text(s, 0.6, 4.95, 12.2, 0.4, "個人からは取らない。場と物で成り立たせる", size=16, font=SERIF)
money = [("場への導入", "科学館・自治体・支援の場に\n体験一式を期間で"), ("持ち帰る物", "カプセル（ブレスレット）・帰還票の台紙"), ("共創", "アーティストと、帰還票の絵柄・会場の演出")]
for i, (h, b) in enumerate(money):
    x = 0.6 + i * 4.2
    panel(s, x, 5.5, 3.95, 1.35)
    text(s, x + 0.25, 5.6, 3.5, 0.45, h, size=15, color=GOLD, bold=True)
    text(s, x + 0.25, 6.05, 3.6, 0.8, b, size=12, color=MUTED)

# 8 Ask -----------------------------------------------------------------------
s = new_slide("08", "Ask", "やりたかったことを、置き去りにしない", """【5:15–5:40】
・わたしはメンターとして、自分を後回しにしてきた人に出会ってきました。MORUNE は、その人たちのやりたかったことを預かっておく場所です。
・まず1回、展示させてください。科学館、プラネタリウム、自治体のイベントの方。
・次に、キャリアや復職を支える場で試させてください。人事の方、キャリア支援をしている方。
・QR から、今ひとつ預けられます。ありがとうございました。""")
asks = [("展示する場", "まず1回。科学館・プラネタリウム・自治体のイベント"),
        ("支える場", "次に。キャリアや復職を支える人事・支援者のもとで試す"),
        ("つくる方", "星空や帰還票の絵柄・映像を一緒に描くアーティスト")]
for i, (h, b) in enumerate(asks):
    y = 1.95 + i * 1.35
    panel(s, 0.6, y, 8.2, 1.15, line=GOLD if i == 0 else None)
    text(s, 0.9, y, 2.3, 1.15, h, size=19, font=SERIF, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 3.2, y, 5.4, 1.15, b, size=14, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 6.15, 8.2, 0.8, "わかっていること：公開2日で、端末12台ほどが願いを預けた（人数ではない）。\nこれから展示と支援の場で、持ち帰った人と、また開いた人を数える", size=12, color=MUTED)
panel(s, 9.4, 1.95, 3.4, 4.6, fill=RGBColor(0xFF, 0xFF, 0xFF))
picture(s, "qr-site.png", 9.65, 2.15, w=2.9)
text(s, 9.4, 5.15, 3.4, 1.3, ["いま、ひとつ預けられます", "morune-25143.morune-25143.workers.dev"], size=11, color=NIGHT, align=PP_ALIGN.CENTER)

prs.save(OUT)
print("saved", OUT)
